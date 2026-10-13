from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

import fitz
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from docx import Document as DocxDocument
from pptx import Presentation

from llama_index.core import (
    Document,
    StorageContext,
    VectorStoreIndex,
    load_index_from_storage,
)
from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.vector_stores import SimpleVectorStore
from llama_index.llms.openai import OpenAI


###############################################################################
# CONFIG
###############################################################################

STORAGE_DIR = Path("storage")
RESULTS_DIR = Path("results")

STORAGE_DIR.mkdir(exist_ok=True)
RESULTS_DIR.mkdir(exist_ok=True)

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
TOP_K = 5


###############################################################################
# PROMPTS
###############################################################################

ASSESSMENT_PROMPT = """
You are an expert reviewer evaluating a scientific, clinical,
regulatory, or research study.

Assessment Criterion:
{criterion}

Evidence:
{evidence}

Scoring Rubric

0-20:
No evidence

21-40:
Weak evidence

41-60:
Partial evidence

61-80:
Strong evidence

81-100:
Comprehensive evidence

Return ONLY valid JSON:

{{
    "score": int,
    "strengths": [str],
    "weaknesses": [str],
    "rationale": str
}}
"""

COMPARISON_PROMPT = """
You are comparing two studies.

Criterion:
{criterion}

Study 1:
{study1_assessment}

Study 2:
{study2_assessment}

Return ONLY valid JSON

{{
    "winner": str,
    "confidence": float,
    "explanation": str
}}
"""


###############################################################################
# DOCUMENT LOADING
###############################################################################


def read_pdf(path: Path) -> str:
    doc = fitz.open(path)

    text = []

    for page in doc:
        text.append(page.get_text())

    return "\n".join(text)


def read_docx(path: Path) -> str:
    doc = DocxDocument(path)

    return "\n".join(
        paragraph.text
        for paragraph in doc.paragraphs
    )


def read_pptx(path: Path) -> str:
    prs = Presentation(path)

    text = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)

    return "\n".join(text)


def load_documents(folder: str, study_name: str):
    documents = []

    folder_path = Path(folder)

    for file in folder_path.rglob("*"):

        if not file.is_file():
            continue

        try:

            suffix = file.suffix.lower()

            content = ""

            if suffix == ".pdf":
                content = read_pdf(file)

            elif suffix == ".docx":
                content = read_docx(file)

            elif suffix == ".pptx":
                content = read_pptx(file)

            if not content.strip():
                continue

            documents.append(
                Document(
                    text=content,
                    metadata={
                        "study": study_name,
                        "filename": file.name,
                        "filetype": suffix,
                    },
                )
            )

        except Exception as ex:
            st.warning(f"Failed reading {file}: {ex}")

    return documents


###############################################################################
# INDEXING
###############################################################################


def build_or_load_index(
    study_name: str,
    docs: list[Document],
):

    storage_path = STORAGE_DIR / study_name

    if storage_path.exists():

        try:

            storage_context = StorageContext.from_defaults(
                persist_dir=str(storage_path)
            )

            return load_index_from_storage(
                storage_context
            )

        except Exception:
            pass

    splitter = SentenceSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )

    nodes = splitter.get_nodes_from_documents(docs)

    vector_store = SimpleVectorStore()

    storage_context = StorageContext.from_defaults(
        vector_store=vector_store
    )

    index = VectorStoreIndex(
        nodes,
        storage_context=storage_context,
    )

    storage_context.persist(
        persist_dir=str(storage_path)
    )

    return index


###############################################################################
# RETRIEVAL
###############################################################################


def retrieve_evidence(
    index: VectorStoreIndex,
    criterion: str,
):

    retriever = index.as_retriever(
        similarity_top_k=TOP_K
    )

    nodes = retriever.retrieve(criterion)

    evidence = []

    combined_text = []

    for node in nodes:

        evidence.append(
            {
                "filename": node.metadata.get(
                    "filename",
                    "unknown",
                ),
                "text": node.text,
                "score": getattr(
                    node,
                    "score",
                    None,
                ),
            }
        )

        combined_text.append(node.text)

    return evidence, "\n\n".join(combined_text)


###############################################################################
# LLM
###############################################################################


def ask_json(llm: OpenAI, prompt: str):

    response = llm.complete(prompt)

    text = str(response)

    start = text.find("{")
    end = text.rfind("}")

    return json.loads(text[start : end + 1])


def assess_study(
    llm,
    criterion,
    evidence_text,
):

    prompt = ASSESSMENT_PROMPT.format(
        criterion=criterion,
        evidence=evidence_text,
    )

    return ask_json(llm, prompt)


def compare_studies(
    llm,
    criterion,
    study1,
    study2,
):

    prompt = COMPARISON_PROMPT.format(
        criterion=criterion,
        study1_assessment=json.dumps(
            study1,
            indent=2,
        ),
        study2_assessment=json.dumps(
            study2,
            indent=2,
        ),
    )

    return ask_json(llm, prompt)


###############################################################################
# DASHBOARD
###############################################################################


def radar_chart(df):

    fig = go.Figure()

    fig.add_trace(
        go.Scatterpolar(
            r=df["study1_score"],
            theta=df["criterion"],
            fill="toself",
            name="Study 1",
        )
    )

    fig.add_trace(
        go.Scatterpolar(
            r=df["study2_score"],
            theta=df["criterion"],
            fill="toself",
            name="Study 2",
        )
    )

    return fig


def bar_chart(df):

    fig = go.Figure()

    fig.add_bar(
        x=df["criterion"],
        y=df["study1_score"],
        name="Study 1",
    )

    fig.add_bar(
        x=df["criterion"],
        y=df["study2_score"],
        name="Study 2",
    )

    return fig


###############################################################################
# APP
###############################################################################


def main():

    st.set_page_config(
        page_title="Comparative Study Assessment",
        layout="wide",
    )

    st.title(
        "Comparative Study Assessment RAG"
    )

    with st.sidebar:

        st.header("Inputs")

        openai_key = st.text_input(
            "OpenAI API Key",
            type="password",
        )

        study1_folder = st.text_input(
            "Study 1 Folder"
        )

        study2_folder = st.text_input(
            "Study 2 Folder"
        )

        criteria_file = st.text_input(
            "Criteria File"
        )

        run_btn = st.button(
            "Run Assessment"
        )

    if not run_btn:
        return

    os.environ["OPENAI_API_KEY"] = openai_key

    llm = OpenAI(
        model="gpt-4o"
    )

    with st.spinner(
        "Loading study documents..."
    ):

        docs1 = load_documents(
            study1_folder,
            "study1",
        )

        docs2 = load_documents(
            study2_folder,
            "study2",
        )

    st.success(
        f"Loaded {len(docs1)} study1 docs "
        f"and {len(docs2)} study2 docs"
    )

    with st.spinner(
        "Building indexes..."
    ):

        idx1 = build_or_load_index(
            "study1",
            docs1,
        )

        idx2 = build_or_load_index(
            "study2",
            docs2,
        )

    criteria = [
        line.strip()
        for line in Path(criteria_file)
        .read_text()
        .splitlines()
        if line.strip()
    ]

    results = []

    progress = st.progress(0)

    for i, criterion in enumerate(criteria):

        ev1, text1 = retrieve_evidence(
            idx1,
            criterion,
        )

        ev2, text2 = retrieve_evidence(
            idx2,
            criterion,
        )

        assessment1 = assess_study(
            llm,
            criterion,
            text1,
        )

        assessment2 = assess_study(
            llm,
            criterion,
            text2,
        )

        comparison = compare_studies(
            llm,
            criterion,
            assessment1,
            assessment2,
        )

        results.append(
            {
                "criterion": criterion,
                "study1_score":
                    assessment1["score"],
                "study2_score":
                    assessment2["score"],
                "study1":
                    assessment1,
                "study2":
                    assessment2,
                "comparison":
                    comparison,
                "evidence1":
                    ev1,
                "evidence2":
                    ev2,
            }
        )

        progress.progress(
            (i + 1) / len(criteria)
        )

    result_df = pd.DataFrame(
        [
            {
                "criterion":
                    r["criterion"],
                "study1_score":
                    r["study1_score"],
                "study2_score":
                    r["study2_score"],
            }
            for r in results
        ]
    )

    overall1 = result_df[
        "study1_score"
    ].mean()

    overall2 = result_df[
        "study2_score"
    ].mean()

    c1, c2, c3 = st.columns(3)

    c1.metric(
        "Overall Study 1",
        f"{overall1:.1f}"
    )

    c2.metric(
        "Overall Study 2",
        f"{overall2:.1f}"
    )

    c3.metric(
        "Winner",
        "Study 1"
        if overall1 > overall2
        else "Study 2"
    )

    st.subheader("Radar Chart")
    st.plotly_chart(
        radar_chart(result_df),
        use_container_width=True,
    )

    st.subheader("Bar Chart")
    st.plotly_chart(
        bar_chart(result_df),
        use_container_width=True,
    )

    st.subheader(
        "Criterion Assessments"
    )

    for r in results:

        with st.expander(
            r["criterion"]
        ):

            st.markdown(
                f"## {r['criterion']}"
            )

            col1, col2 = st.columns(2)

            with col1:

                st.metric(
                    "Study 1 Score",
                    r["study1_score"],
                )

                st.write(
                    r["study1"]["rationale"]
                )

                st.write("### Evidence")

                for ev in r["evidence1"]:
                    st.write(
                        f"Source: {ev['filename']}"
                    )
                    st.write(ev["text"][:800])

            with col2:

                st.metric(
                    "Study 2 Score",
                    r["study2_score"],
                )

                st.write(
                    r["study2"]["rationale"]
                )

                st.write("### Evidence")

                for ev in r["evidence2"]:
                    st.write(
                        f"Source: {ev['filename']}"
                    )
                    st.write(ev["text"][:800])

            st.write("### Comparison")

            st.json(
                r["comparison"]
            )

    json_path = (
        RESULTS_DIR
        / "assessment_results.json"
    )

    csv_path = (
        RESULTS_DIR
        / "assessment_results.csv"
    )

    json.dump(
        results,
        open(json_path, "w"),
        indent=2,
    )

    result_df.to_csv(
        csv_path,
        index=False,
    )

    st.success(
        f"Saved results to {json_path}"
    )


if __name__ == "__main__":
    main()