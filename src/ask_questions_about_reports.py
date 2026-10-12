import os
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation

from llama_index.core import (
    VectorStoreIndex,
    Document as LlamaDocument,
    Settings
)

from llama_index.embeddings.openai import OpenAIEmbedding
from llama_index.llms.openai import OpenAI


# ---------------------------------------------------
# EXTRACTORS
# ---------------------------------------------------

def read_pdf(path):
    reader = PdfReader(path)
    text = []

    for page in reader.pages:
        t = page.extract_text()
        if t:
            text.append(t)

    return "\n".join(text)


def read_docx(path):
    doc = Document(path)

    return "\n".join(
        p.text
        for p in doc.paragraphs
        if p.text.strip()
    )


def read_pptx(path):
    prs = Presentation(path)

    text = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)

    return "\n".join(text)


# ---------------------------------------------------
# LOAD FOLDER
# ---------------------------------------------------

def load_documents(folder):

    documents = []

    for root, _, files in os.walk(folder):

        for file in files:

            full_path = os.path.join(root, file)

            try:

                if file.endswith(".pdf"):
                    content = read_pdf(full_path)

                elif file.endswith(".docx"):
                    content = read_docx(full_path)

                elif file.endswith(".pptx"):
                    content = read_pptx(full_path)

                else:
                    continue

                doc = LlamaDocument(
                    text=content,
                    metadata={
                        "source": full_path
                    }
                )

                documents.append(doc)

                print(f"Loaded {full_path}")

            except Exception as e:
                print(e)

    return documents


# ---------------------------------------------------
# MAIN
# ---------------------------------------------------

folder = "documents"

Settings.embed_model = OpenAIEmbedding(
    model="text-embedding-3-small"
)

Settings.llm = OpenAI(
    model="gpt-4o-mini",
    temperature=0,
)

docs = load_documents(folder)

index = VectorStoreIndex.from_documents(docs)

query_engine = index.as_query_engine(
    similarity_top_k=5
)

while True:

    q = input("\nQuestion: ")

    if q.lower() == "exit":
        break

    response = query_engine.query(q)

    print(response)