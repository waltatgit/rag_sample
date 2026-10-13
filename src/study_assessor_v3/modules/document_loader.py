"""Module responsible for loading and parsing PDF, DOCX, and PPTX documents."""

import os
from typing import List
import fitz  # PyMuPDF
import docx
import pptx
from llama_index.core.schema import Document
from modules.utils import setup_logger

logger = setup_logger()


class DocumentLoader:
    """Parses PDF, DOCX, and PPTX files into LlamaIndex Document instances."""

    def __init__(self, folder_path: str, study_label: str):
        self.folder_path = folder_path
        self.study_label = study_label

    def load_documents(self) -> List[Document]:
        """Scans folder and returns list of LlamaIndex Document objects with metadata."""
        if not os.path.exists(self.folder_path):
            raise FileNotFoundError(f"Folder directory does not exist: {self.folder_path}")

        documents: List[Document] = []
        for root, _, files in os.walk(self.folder_path):
            for filename in files:
                filepath = os.path.join(root, filename)
                ext = os.path.splitext(filename)[1].lower()
                text = ""

                try:
                    if ext == ".pdf":
                        text = self._parse_pdf(filepath)
                    elif ext == ".docx":
                        text = self._parse_docx(filepath)
                    elif ext == ".pptx":
                        text = self._parse_pptx(filepath)
                    else:
                        continue  # Skip unsupported extensions

                    if text.strip():
                        doc = Document(
                            text=text,
                            metadata={
                                "study": self.study_label,
                                "filename": filename,
                                "filetype": ext.lstrip(".")
                            }
                        )
                        documents.append(doc)
                        logger.info(f"Loaded document: {filename} for {self.study_label}")
                except Exception as e:
                    logger.error(f"Error reading file {filepath}: {e}")

        return documents

    @staticmethod
    def _parse_pdf(filepath: str) -> str:
        text = []
        with fitz.open(filepath) as doc:
            for page in doc:
                text.append(page.get_text())
        return "\n".join(text)

    @staticmethod
    def _parse_docx(filepath: str) -> str:
        doc = docx.Document(filepath)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

    @staticmethod
    def _parse_pptx(filepath: str) -> str:
        prs = pptx.Presentation(filepath)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)